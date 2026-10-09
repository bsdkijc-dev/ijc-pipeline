"""Load the IJC workbook from Google Sheets (if credentials exist) or the local xlsx."""
import json, os, re
from datetime import datetime, timedelta

TABLES = {  # sheet name -> header cell that identifies the header row
    "Master_Knowledge_Objects": "Knowledge ID",
    "Evidence_Provenance": "Evidence ID",
    "Atomic_Claims_Graph": "Claim ID",
    "GEO_QA_Triage": "Query ID",
    "Verification_QualityGate": "Gate Check ID",
}


def _rows_to_dicts(rows, key):
    rows = [[("" if c is None else c) for c in r] for r in rows]
    hdr_i = next(i for i, r in enumerate(rows) if key in r)
    hdr = rows[hdr_i]
    width = len(hdr)
    rows = [r + [""] * (width - len(r)) for r in rows]
    cols = [j for j, h in enumerate(hdr) if str(h).strip()]
    out = []
    for r in rows[hdr_i + 1:]:
        if not str(r[hdr.index(key)]).strip():
            continue
        out.append({str(hdr[j]).strip(): r[j] if j < len(r) else "" for j in cols})
    return out


def load():
    sheet_id, sa = os.getenv("GOOGLE_SHEETS_ID"), os.getenv("GCP_SA_KEY")
    data = {}
    if sheet_id and sa:
        import gspread
        gc = gspread.service_account_from_dict(json.loads(sa))
        sh = gc.open_by_key(sheet_id)
        for name, key in TABLES.items():
            data[name] = _rows_to_dicts(sh.worksheet(name).get_all_values(), key)
        print("source: Google Sheets")
    else:
        from openpyxl import load_workbook
        wb = load_workbook("data/ijc.xlsx", data_only=True, read_only=True)
        for name, key in TABLES.items():
            data[name] = _rows_to_dicts(wb[name].iter_rows(values_only=True), key)
        print("source: data/ijc.xlsx")
    return data


def excel_date(v):
    if isinstance(v, (int, float)) and v > 20000:
        return (datetime(1899, 12, 30) + timedelta(days=v)).strftime("%Y-%m-%d")
    return str(v)


def slug(s):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(s)).strip("_")
